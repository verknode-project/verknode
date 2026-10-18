"""More IO nodes: Excel, native tables, JSON, XML, lines, fixed width, models, images, file handling."""
from __future__ import annotations

import datetime as dt
import fnmatch
import gzip
import json
import os
import pickle
import shutil
import tarfile
import zipfile

import polars as pl

from verknode.node import (
    F,
    FLOWVAR,
    FLOWVARIABLE,
    IMAGE,
    InvalidSettingsError,
    MODEL,
    ModelPortObject,
    NodeModel,
    Port,
    TABLE,
    require_column,
)
from verknode.node import CELL_DTYPES, CELL_TYPES, dtype_kind, json_value, parse_cell
from .io import PATH_HELP, _existing

IF_EXISTS = ["Overwrite", "Fail"]


def _check_target(ctx, path: str, if_exists: str) -> str:
    p = ctx.resolve_path(path)
    if os.path.exists(p) and if_exists == "Fail":
        raise InvalidSettingsError(f"File exists and must not be overwritten: {p}")
    return p


def _mkparents(p: str):
    os.makedirs(os.path.dirname(p) or ".", exist_ok=True)


def _infer_strings(name: str, values: list) -> pl.Series:
    """String values -> Boolean / Int64 / Float64 column where all non-missing values parse."""
    s = pl.Series(name, values, dtype=pl.String)
    present = s.len() - s.null_count()
    if present:
        low = s.str.strip_chars().str.to_lowercase()
        if low.drop_nulls().is_in(["true", "false"]).all():
            return (low == "true").alias(name)
        for dtype in (pl.Int64, pl.Float64):
            cast = s.str.strip_chars().cast(dtype, strict=False)
            if cast.null_count() == s.null_count():
                return cast
    return s


def _python_column(name: str, values: list) -> pl.Series:
    """Column from JSON-like Python values: bool/int/float/str kept, nested values as JSON text."""
    present = [v for v in values if v is not None]
    if present and all(isinstance(v, bool) for v in present):
        return pl.Series(name, values, dtype=pl.Boolean)
    if present and all(isinstance(v, int) and not isinstance(v, bool) for v in present):
        return pl.Series(name, values, dtype=pl.Int64)
    if present and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in present):
        return pl.Series(name, [None if v is None else float(v) for v in values], dtype=pl.Float64)
    return pl.Series(name, [None if v is None else v if isinstance(v, str) else json.dumps(v)
                            for v in values], dtype=pl.String)


def _records_to_df(records: list) -> pl.DataFrame:
    columns: dict[str, list] = {}
    for i, rec in enumerate(records):
        if not isinstance(rec, dict):
            rec = {"Value": rec}
        for k in rec:
            columns.setdefault(str(k), [None] * i)
        for k, col in columns.items():
            col.append(rec.get(k))
    return pl.DataFrame([_python_column(k, v) for k, v in columns.items()])


def _to_text_columns(df: pl.DataFrame) -> pl.DataFrame:
    """Molecules -> SMILES; lists, structs and binary -> JSON text (for Excel / JSON output)."""
    if any(dtype_kind(t) == "molecule" for t in df.schema.values()):
        from verknode.node import chem
        molecules_to_smiles = chem.molecules_to_smiles

        df = molecules_to_smiles(df)
    exprs = []
    for name, t in df.schema.items():
        if dtype_kind(t) in ("list", "struct", "binary", "other"):
            exprs.append(pl.Series(name, [None if v is None else json.dumps(json_value(v)) for v in df[name].to_list()],
                                   dtype=pl.String))
    return df.with_columns(exprs) if exprs else df


# ----------------------------------------------------------------------
# Read
# ----------------------------------------------------------------------
def _excel_sheet_kwargs(s) -> dict:
    if s["sheet_mode"] == "Sheet by name":
        if not s["sheet_name"]:
            raise InvalidSettingsError("No sheet name specified")
        return {"sheet_name": s["sheet_name"]}
    if s["sheet_mode"] == "Sheet by position":
        return {"sheet_id": max(1, s["sheet_index"])}
    return {"sheet_id": 1}


class ExcelReader(NodeModel):
    factory = "verknode.io.ExcelReader"
    name = "Excel Reader"
    category = "IO/Read"
    node_type = "Source"
    icon = "XLS"
    short_description = "Reads a sheet of an Excel file (.xlsx, .xls, .xlsb, .ods)."
    description = (
        "Reads one sheet of a spreadsheet into a table. Column types are detected from the cell values. "
        "Use <i>Columns</i> (e.g. <code>A:D</code> or <code>A,C,E:F</code>) and <i>Skip rows</i> to read only a range."
    )
    out_ports = [Port(TABLE, "Output table", "Contents of the sheet")]
    fields = [
        F("path", "string", "File", "", PATH_HELP, path=True),
        F("sheet_mode", "enum", "Sheet", "First sheet", options=["First sheet", "Sheet by name", "Sheet by position"]),
        F("sheet_name", "string", "Sheet name", "", show_if={"sheet_mode": "Sheet by name"}),
        F("sheet_index", "int", "Sheet position", 1, "1 = first sheet.", show_if={"sheet_mode": "Sheet by position"}),
        F("has_header", "bool", "Table contains column names", True),
        F("skip_rows", "int", "Skip first rows", 0, "Rows skipped before the header / first data row."),
        F("columns", "string", "Columns", "", "Excel column letters, e.g. A:D (empty: all)."),
        F("n_rows", "int", "Maximum rows (0 = all)", 0),
    ]

    def execute(self, inputs, s, ctx):
        p = _existing(ctx, s["path"])
        read_options = {"skip_rows": max(0, s["skip_rows"])}
        if s["columns"]:
            read_options["use_columns"] = s["columns"].replace(" ", "")
        if s["n_rows"] > 0:
            read_options["n_rows"] = s["n_rows"]
        import fastexcel  # noqa: F401 - (polars reads Excel files with it)

        try:
            df = pl.read_excel(p, **_excel_sheet_kwargs(s), has_header=s["has_header"],
                               read_options=read_options, raise_if_empty=False)
        except InvalidSettingsError:
            raise
        except Exception as e:  # noqa: BLE001 - surface reader errors as configuration problems
            raise InvalidSettingsError(f"Cannot read Excel file: {e}") from None
        return [df]


class ReadExcelSheetNames(NodeModel):
    factory = "verknode.io.ReadExcelSheetNames"
    name = "Read Excel Sheet Names"
    category = "IO/Read"
    node_type = "Source"
    icon = "XLS"
    short_description = "Lists the sheet names of an Excel file."
    description = "Outputs one row per sheet with the sheet name and the file path, e.g. to loop over all sheets."
    out_ports = [Port(TABLE, "Sheet names")]
    fields = [F("path", "string", "File", "", PATH_HELP, path=True)]

    def execute(self, inputs, s, ctx):
        import fastexcel

        p = _existing(ctx, s["path"])
        names = fastexcel.read_excel(p).sheet_names
        return [pl.DataFrame({"Path": pl.Series([p] * len(names), dtype=pl.String),
                              "Sheet": pl.Series(names, dtype=pl.String)})]


class TableReader(NodeModel):
    factory = "verknode.io.TableReader"
    name = "Table Reader"
    category = "IO/Read"
    node_type = "Source"
    icon = "TBL"
    short_description = "Reads a table written by the Table Writer."
    description = "Reads the native table format (Apache Arrow IPC, extension <code>.table</code>), keeping all column types including molecules."
    out_ports = [Port(TABLE, "Read table")]
    fields = [F("path", "string", "File", "", PATH_HELP, path=True)]

    def configure(self, in_specs, s, ctx):
        return [pl.read_ipc_schema(_existing(ctx, s["path"]))]

    def execute(self, inputs, s, ctx):
        return [pl.read_ipc(_existing(ctx, s["path"]))]


def _json_path(data, path: str):
    if not path or path in ("$", "$."):
        return data
    from jsonpath_ng.ext import parse as jsonpath

    try:
        matches = [m.value for m in jsonpath(path).find(data)]
    except Exception as e:  # noqa: BLE001
        raise InvalidSettingsError(f"Invalid JSONPath {path!r}: {e}") from None
    if len(matches) == 1 and isinstance(matches[0], list):
        return matches[0]
    return matches


class JSONReader(NodeModel):
    factory = "verknode.io.JSONReader"
    name = "JSON Reader"
    category = "IO/Read"
    node_type = "Source"
    icon = "{ }"
    short_description = "Reads a JSON or JSON Lines file into a table."
    description = (
        "Reads a JSON file. A list of objects becomes one row per object and one column per key; nested "
        "values are kept as JSON text. Use a JSONPath (e.g. <code>$.data.items</code>) to select the list. "
        "JSON Lines files contain one JSON value per line."
    )
    out_ports = [Port(TABLE, "JSON table")]
    fields = [
        F("path", "string", "File", "", PATH_HELP, path=True),
        F("format", "enum", "Format", "JSON", options=["JSON", "JSON Lines"]),
        F("json_path", "string", "Select with JSONPath", "$", "Path to the list of records, $ = whole document.",
          show_if={"format": "JSON"}),
    ]

    def execute(self, inputs, s, ctx):
        p = _existing(ctx, s["path"])
        opener = gzip.open if p.endswith(".gz") else open
        with opener(p, "rt", encoding="utf-8") as fh:
            try:
                if s["format"] == "JSON Lines":
                    data = [json.loads(line) for line in fh if line.strip()]
                else:
                    data = _json_path(json.load(fh), s["json_path"])
            except json.JSONDecodeError as e:
                raise InvalidSettingsError(f"Invalid JSON: {e}") from None
        if not isinstance(data, list):
            data = [data]
        return [_records_to_df(data)]


def _local(tag) -> str:
    return tag.split("}", 1)[-1] if isinstance(tag, str) else str(tag)


class XMLReader(NodeModel):
    factory = "verknode.io.XMLReader"
    name = "XML Reader"
    category = "IO/Read"
    node_type = "Source"
    icon = "XML"
    short_description = "Reads elements of an XML file into a table."
    description = (
        "Selects row elements with an XPath (default: all children of the root element). Each attribute "
        "(column <code>@name</code>) and each child element's text becomes a column; types are detected. "
        "Namespaces are ignored in column names; use <code>local-name()</code> in XPaths for namespaced documents."
    )
    out_ports = [Port(TABLE, "XML table")]
    fields = [
        F("path", "string", "File", "", PATH_HELP, path=True),
        F("xpath", "string", "Row XPath", "/*/*", "XPath selecting one element per row."),
        F("attributes", "bool", "Read attributes", True),
        F("children", "bool", "Read child elements", True),
        F("text", "bool", "Read element text (column 'text')", False),
    ]

    def execute(self, inputs, s, ctx):
        from lxml import etree

        p = _existing(ctx, s["path"])
        try:
            tree = etree.parse(p, etree.XMLParser(resolve_entities=False, no_network=True))
            elements = tree.xpath(s["xpath"] or "/*/*")
        except (etree.XMLSyntaxError, etree.XPathError) as e:
            raise InvalidSettingsError(f"Cannot read XML: {e}") from None
        records = []
        for el in elements:
            if not isinstance(el, etree._Element):
                records.append({"value": str(el)})
                continue
            rec = {}
            if s["attributes"]:
                rec.update({f"@{_local(k)}": v for k, v in el.attrib.items()})
            if s["children"]:
                for child in el:
                    if isinstance(child.tag, str):
                        rec.setdefault(_local(child.tag), (child.text or "").strip())
            if s["text"]:
                rec["text"] = (el.text or "").strip()
            records.append(rec)
        columns: dict[str, list] = {}
        for i, rec in enumerate(records):
            for k in rec:
                columns.setdefault(k, [None] * i)
            for k, col in columns.items():
                col.append(rec.get(k))
        return [pl.DataFrame([_infer_strings(k, v) for k, v in columns.items()])]


class LineReader(NodeModel):
    factory = "verknode.io.LineReader"
    name = "Line Reader"
    category = "IO/Read"
    node_type = "Source"
    icon = "≡"
    short_description = "Reads a text file line by line."
    description = "Reads each line of a text file into one row of a String column."
    out_ports = [Port(TABLE, "Lines")]
    fields = [
        F("path", "string", "File", "", PATH_HELP, path=True),
        F("column", "string", "Column name", "Column"),
        F("skip_empty", "bool", "Skip empty lines", False),
        F("use_header", "bool", "Use first line as column name", False),
        F("limit", "int", "Maximum lines (0 = all)", 0),
        F("encoding", "string", "Encoding", "utf-8"),
    ]

    def execute(self, inputs, s, ctx):
        p = _existing(ctx, s["path"])
        opener = gzip.open if p.endswith(".gz") else open
        lines = []
        name = s["column"] or "Column"
        try:
            with opener(p, "rt", encoding=s["encoding"] or "utf-8", errors="replace") as fh:
                for i, line in enumerate(fh):
                    line = line.rstrip("\r\n")
                    if i == 0 and s["use_header"]:
                        name = line or name
                        continue
                    if s["skip_empty"] and not line.strip():
                        continue
                    lines.append(line)
                    if s["limit"] and len(lines) >= s["limit"]:
                        break
        except LookupError:
            raise InvalidSettingsError(f"Unknown encoding: {s['encoding']}") from None
        return [pl.DataFrame({name: pl.Series(lines, dtype=pl.String)})]


class FixedWidthFileReader(NodeModel):
    factory = "verknode.io.FixedWidthFileReader"
    name = "Fixed Width File Reader"
    category = "IO/Read"
    node_type = "Source"
    icon = "|||"
    short_description = "Reads a text file with fixed width columns."
    description = (
        "Each column is defined by its name, start position (0-based character offset), width and type. "
        "Values are stripped of surrounding spaces; empty values are missing."
    )
    out_ports = [Port(TABLE, "Output table")]
    fields = [
        F("path", "string", "File", "", PATH_HELP, path=True),
        F("columns", "records", "Columns", [{"name": "column1", "start": 0, "width": 10, "type": "String"}], fields=[
            F("name", "string", "Name", ""),
            F("start", "int", "Start", 0),
            F("width", "int", "Width", 10),
            F("type", "enum", "Type", "String", options=CELL_TYPES),
        ]),
        F("skip_lines", "int", "Skip first lines", 0, "e.g. 1 to skip a header line."),
        F("skip_empty", "bool", "Skip empty lines", True),
    ]

    def execute(self, inputs, s, ctx):
        p = _existing(ctx, s["path"])
        cols = [c for c in s["columns"] if c.get("name")]
        if not cols:
            raise InvalidSettingsError("No columns defined")
        if len({c["name"] for c in cols}) != len(cols):
            raise InvalidSettingsError("Column names must be unique")
        values = {c["name"]: [] for c in cols}
        with open(p, encoding="utf-8", errors="replace") as fh:
            for i, line in enumerate(fh):
                line = line.rstrip("\r\n")
                if i < s["skip_lines"] or (s["skip_empty"] and not line.strip()):
                    continue
                for c in cols:
                    raw = line[c["start"]:c["start"] + c["width"]].strip()
                    try:
                        values[c["name"]].append(parse_cell(raw, c["type"]) if raw else None)
                    except (TypeError, ValueError):
                        raise InvalidSettingsError(f'Line {i + 1}, column "{c["name"]}": cannot read {raw!r} as {c["type"]}') from None
        return [pl.DataFrame({c["name"]: pl.Series(values[c["name"]], dtype=CELL_DTYPES[c["type"]]) for c in cols})]


def _load_model(path: str) -> ModelPortObject:
    with open(path, "rb") as fh:
        try:
            obj = pickle.load(fh)  # noqa: S301 - model files are trusted like workflows
        except Exception as e:  # noqa: BLE001
            raise InvalidSettingsError(f"Cannot read model file: {e}") from None
    if not isinstance(obj, ModelPortObject):
        raise InvalidSettingsError("File does not contain a model written by the Model Writer")
    return obj


class ModelReader(NodeModel):
    factory = "verknode.io.ModelReader"
    name = "Model Reader"
    category = "IO/Read"
    node_type = "Source"
    icon = "MDL"
    short_description = "Reads a model written by the Model Writer."
    description = "Reads a trained model (e.g. from a learner node) saved with the Model Writer."
    out_ports = [Port(MODEL, "Model")]
    fields = [F("path", "string", "File", "", PATH_HELP, path=True)]

    def configure(self, in_specs, s, ctx):
        return [_load_model(_existing(ctx, s["path"])).spec()]

    def execute(self, inputs, s, ctx):
        return [_load_model(_existing(ctx, s["path"]))]


# ----------------------------------------------------------------------
# Write
# ----------------------------------------------------------------------
class _FileSink(NodeModel):
    category = "IO/Write"
    node_type = "Sink"
    out_ports = []

    def configure(self, in_specs, s, ctx):
        _check_target(ctx, s["path"], s.get("if_exists", "Overwrite"))
        return []


class ExcelWriter(_FileSink):
    factory = "verknode.io.ExcelWriter"
    name = "Excel Writer"
    icon = "XLS"
    short_description = "Writes a table to a sheet of an Excel file."
    description = (
        "Writes the input table to an .xlsx file. <i>Append sheet</i> adds (or replaces) the sheet in an "
        "existing workbook. Molecules are written as SMILES, collections as JSON text."
    )
    in_ports = [Port(TABLE, "Input table")]
    fields = [
        F("path", "string", "File", "", PATH_HELP, path=True),
        F("sheet", "string", "Sheet name", "default"),
        F("include_header", "bool", "Write column header", True),
        F("if_exists", "enum", "If file exists", "Overwrite", options=["Overwrite", "Append sheet", "Fail"]),
        F("autofit", "bool", "Autosize columns", True),
    ]

    def execute(self, inputs, s, ctx):
        df = _to_text_columns(inputs[0])
        p = _check_target(ctx, s["path"], s["if_exists"])
        sheet = s["sheet"] or "default"
        _mkparents(p)
        if s["if_exists"] == "Append sheet" and os.path.exists(p):
            from openpyxl import load_workbook

            wb = load_workbook(p)
            if sheet in wb.sheetnames:
                del wb[sheet]
            ws = wb.create_sheet(sheet)
            if s["include_header"]:
                ws.append(df.columns)
            for row in df.iter_rows():
                ws.append([v.isoformat() if isinstance(v, dt.timedelta) else v for v in row])
            wb.save(p)
        else:
            import xlsxwriter  # noqa: F401 - (polars writes Excel files with it)

            df.write_excel(p, worksheet=sheet, include_header=s["include_header"], autofit=s["autofit"],
                           autofilter=False)
        ctx.log(f"Wrote {df.height} rows to sheet {sheet!r} of {p}")
        return []


class TableWriter(_FileSink):
    factory = "verknode.io.TableWriter"
    name = "Table Writer"
    icon = "TBL"
    short_description = "Writes a table in the native format."
    description = "Writes the table in the native format (Apache Arrow IPC, extension <code>.table</code>), keeping all column types."
    in_ports = [Port(TABLE, "Input table")]
    fields = [
        F("path", "string", "File", "", PATH_HELP, path=True),
        F("compression", "enum", "Compression", "zstd", options=["zstd", "lz4", "uncompressed"]),
        F("if_exists", "enum", "If file exists", "Overwrite", options=IF_EXISTS),
    ]

    def execute(self, inputs, s, ctx):
        p = _check_target(ctx, s["path"], s["if_exists"])
        _mkparents(p)
        inputs[0].write_ipc(p, compression=s["compression"])
        ctx.log(f"Wrote {inputs[0].height} rows to {p}")
        return []


class JSONWriter(_FileSink):
    factory = "verknode.io.JSONWriter"
    name = "JSON Writer"
    icon = "{ }"
    short_description = "Writes a table to a JSON or JSON Lines file."
    description = (
        "Writes rows as an array of objects (<i>Records</i>), as an object of column arrays (<i>Columns</i>) or as "
        "one object per line (<i>JSON Lines</i>). Dates are written in ISO 8601, molecules as SMILES."
    )
    in_ports = [Port(TABLE, "Input table")]
    fields = [
        F("path", "string", "File", "", PATH_HELP, path=True),
        F("orientation", "enum", "Orientation", "Records", options=["Records", "Columns", "JSON Lines"]),
        F("pretty", "bool", "Pretty print", True, show_if={"orientation": ["Records", "Columns"]}),
        F("if_exists", "enum", "If file exists", "Overwrite", options=IF_EXISTS),
    ]

    def execute(self, inputs, s, ctx):
        df = inputs[0]
        if any(dtype_kind(t) == "molecule" for t in df.schema.values()):
            from verknode.node import chem
            molecules_to_smiles = chem.molecules_to_smiles

            df = molecules_to_smiles(df)
        p = _check_target(ctx, s["path"], s["if_exists"])
        _mkparents(p)
        indent = 2 if s["pretty"] else None
        with open(p, "w", encoding="utf-8") as fh:
            if s["orientation"] == "JSON Lines":
                for row in df.iter_rows(named=True):
                    fh.write(json.dumps(json_value(row), ensure_ascii=False) + "\n")
            elif s["orientation"] == "Columns":
                json.dump({c: json_value(df[c].to_list()) for c in df.columns}, fh, indent=indent, ensure_ascii=False)
            else:
                json.dump([json_value(r) for r in df.iter_rows(named=True)], fh, indent=indent, ensure_ascii=False)
        ctx.log(f"Wrote {df.height} rows to {p}")
        return []


class ModelWriter(_FileSink):
    factory = "verknode.io.ModelWriter"
    name = "Model Writer"
    icon = "MDL"
    short_description = "Writes a model to a file."
    description = "Saves a trained model so it can be read again with the Model Reader."
    in_ports = [Port(MODEL, "Model")]
    fields = [
        F("path", "string", "File", "", PATH_HELP + " Suggested extension: .model", path=True),
        F("if_exists", "enum", "If file exists", "Overwrite", options=IF_EXISTS),
    ]

    def execute(self, inputs, s, ctx):
        p = _check_target(ctx, s["path"], s["if_exists"])
        _mkparents(p)
        with open(p, "wb") as fh:
            pickle.dump(inputs[0], fh)
        ctx.log(f"Wrote model to {p}")
        return []


class ImageWriterPort(_FileSink):
    factory = "verknode.io.ImageWriterPort"
    name = "Image Writer (Port)"
    icon = "IMG"
    short_description = "Writes an image port to a file."
    description = "Writes the image from an image port (e.g. of a view node) to a file (.svg or .png, matching the image type)."
    in_ports = [Port(IMAGE, "Image")]
    fields = [
        F("path", "string", "File", "", PATH_HELP, path=True),
        F("if_exists", "enum", "If file exists", "Overwrite", options=IF_EXISTS),
    ]

    def execute(self, inputs, s, ctx):
        img = inputs[0]
        p = _check_target(ctx, s["path"], s["if_exists"])
        _mkparents(p)
        with open(p, "wb") as fh:
            fh.write(img.data)
        ctx.log(f"Wrote {img.mime} image to {p}")
        return []


# ----------------------------------------------------------------------
# File handling
# ----------------------------------------------------------------------
FH = "IO/File Handling"

LIST_SCHEMA = pl.Schema({
    "Path": pl.String,
    "Name": pl.String,
    "Extension": pl.String,
    "Size": pl.Int64,
    "Last modified": pl.Datetime("us"),
    "Is directory": pl.Boolean,
})


def _folder(ctx, path: str) -> str:
    p = ctx.resolve_path(path)
    if not os.path.isdir(p):
        raise InvalidSettingsError(f"Folder does not exist: {p}")
    return p


class ListFilesFolders(NodeModel):
    factory = "verknode.io.ListFilesFolders"
    name = "List Files/Folders"
    category = FH
    node_type = "Source"
    icon = "📁"
    short_description = "Lists files and folders of a folder."
    description = (
        "Lists the content of a folder, optionally including subfolders. Filter files by extension "
        "(e.g. <code>csv;txt</code>) or a wildcard pattern (e.g. <code>sales_*.xlsx</code>)."
    )
    out_ports = [Port(TABLE, "File/folder list")]
    fields = [
        F("folder", "string", "Folder", ".", PATH_HELP, path=True),
        F("recursive", "bool", "Include subfolders", False),
        F("mode", "enum", "List", "Files", options=["Files", "Folders", "Files and folders"]),
        F("filter", "enum", "Filter", "None", options=["None", "File extensions", "Wildcard"]),
        F("extensions", "string", "Extensions", "csv", "Separated by ;", show_if={"filter": "File extensions"}),
        F("pattern", "string", "Wildcard pattern", "*", show_if={"filter": "Wildcard"}),
        F("case_sensitive", "bool", "Case sensitive filter", False, show_if={"filter": ["File extensions", "Wildcard"]}),
        F("hidden", "bool", "Include hidden files", False),
    ]

    def configure(self, in_specs, s, ctx):
        _folder(ctx, s["folder"])
        return [LIST_SCHEMA]

    def _matches(self, name, s):
        if s["filter"] == "None":
            return True
        target = name if s["case_sensitive"] else name.lower()
        if s["filter"] == "Wildcard":
            pat = s["pattern"] if s["case_sensitive"] else (s["pattern"] or "*").lower()
            return fnmatch.fnmatchcase(target, pat or "*")
        exts = [e.strip().lstrip(".") for e in (s["extensions"] or "").split(";") if e.strip()]
        if not s["case_sensitive"]:
            exts = [e.lower() for e in exts]
        return any(target.endswith("." + e) for e in exts)

    def execute(self, inputs, s, ctx):
        root = _folder(ctx, s["folder"])
        rows = []
        for dirpath, dirnames, filenames in os.walk(root):
            ctx.check_canceled()
            if not s["hidden"]:
                dirnames[:] = [d for d in dirnames if not d.startswith(".")]
                filenames = [f for f in filenames if not f.startswith(".")]
            entries = []
            if s["mode"] in ("Folders", "Files and folders"):
                entries += [(d, True) for d in sorted(dirnames)]
            if s["mode"] in ("Files", "Files and folders"):
                entries += [(f, False) for f in sorted(filenames) if self._matches(f, s)]
            for name, is_dir in entries:
                full = os.path.join(dirpath, name)
                try:
                    st = os.stat(full)
                except OSError:
                    continue
                rows.append((full, name, "" if is_dir else os.path.splitext(name)[1].lstrip("."),
                             None if is_dir else st.st_size, dt.datetime.fromtimestamp(st.st_mtime), is_dir))
            if not s["recursive"]:
                break
        cols = list(zip(*rows)) if rows else [[]] * len(LIST_SCHEMA)
        return [pl.DataFrame({n: pl.Series(n, list(v), dtype=t) for (n, t), v in zip(LIST_SCHEMA.items(), cols)})]


class CreateFolder(NodeModel):
    factory = "verknode.io.CreateFolder"
    name = "Create Folder"
    category = FH
    node_type = "Other"
    icon = "📁+"
    short_description = "Creates a folder and exposes its path as flow variable."
    description = "Creates a folder (including missing parent folders) and pushes its absolute path as a String flow variable."
    out_ports = [Port(FLOWVARIABLE, "Variables connection", "Carries the folder path variable")]
    fields = [
        F("parent", "string", "Parent folder", ".", PATH_HELP, path=True),
        F("folder", "string", "Folder name", "new_folder", "May contain subfolders, e.g. results/run1."),
        F("if_exists", "enum", "If folder exists", "Ignore", options=["Ignore", "Fail"]),
        F("variable", "string", "Path variable name", "folder_path"),
    ]

    def _path(self, s, ctx):
        if not s["folder"]:
            raise InvalidSettingsError("No folder name specified")
        if not s["variable"]:
            raise InvalidSettingsError("No variable name specified")
        return os.path.normpath(os.path.join(ctx.resolve_path(s["parent"] or "."), s["folder"]))

    def configure(self, in_specs, s, ctx):
        ctx.push_variable(s["variable"], self._path(s, ctx), "STRING")
        return [FLOWVAR]

    def execute(self, inputs, s, ctx):
        p = self._path(s, ctx)
        if os.path.exists(p):
            if s["if_exists"] == "Fail":
                raise InvalidSettingsError(f"Folder already exists: {p}")
            if not os.path.isdir(p):
                raise InvalidSettingsError(f"A file with this name exists: {p}")
        os.makedirs(p, exist_ok=True)
        ctx.push_variable(s["variable"], p, "STRING")
        return [FLOWVAR]


def _path_column_schema(spec, extra: dict):
    out = dict(spec)
    for k, v in extra.items():
        if k in out:
            raise InvalidSettingsError(f'Column "{k}" already exists')
        out[k] = v
    return pl.Schema(out)


class DeleteFilesFolders(NodeModel):
    factory = "verknode.io.DeleteFilesFolders"
    name = "Delete Files/Folders"
    category = FH
    node_type = "Sink"
    icon = "🗑"
    short_description = "Deletes the files and folders listed in a path column."
    description = "Deletes each file or folder in the path column and appends a <i>Status</i> column (deleted / not found)."
    in_ports = [Port(TABLE, "Paths")]
    out_ports = [Port(TABLE, "Deletion status")]
    fields = [
        F("column", "column", "Path column", None, kinds=["string"]),
        F("recursive", "bool", "Delete folders with their content", False),
        F("fail_missing", "bool", "Fail if a file does not exist", False),
    ]

    def configure(self, in_specs, s, ctx):
        require_column(s, "column", in_specs[0], "path column")
        return [_path_column_schema(in_specs[0], {"Status": pl.String})]

    def execute(self, inputs, s, ctx):
        df = inputs[0]
        col = require_column(s, "column", df.schema, "path column")
        status = []
        for v in df[col].to_list():
            ctx.check_canceled()
            p = ctx.resolve_path(v) if v else None
            if p is None or not os.path.lexists(p):
                if s["fail_missing"]:
                    raise InvalidSettingsError(f"File does not exist: {v}")
                status.append("not found")
            elif os.path.isdir(p) and not os.path.islink(p):
                if s["recursive"]:
                    shutil.rmtree(p)
                else:
                    os.rmdir(p)
                status.append("deleted")
            else:
                os.remove(p)
                status.append("deleted")
        return [df.with_columns(pl.Series("Status", status, dtype=pl.String))]


TRANSFER_SCHEMA = pl.Schema({"Source Path": pl.String, "Destination Path": pl.String, "Status": pl.String})


class TransferFiles(NodeModel):
    factory = "verknode.io.TransferFiles"
    name = "Transfer Files"
    category = FH
    node_type = "Other"
    icon = "⇉"
    short_description = "Copies or moves files and folders to a destination folder."
    description = "Copies or moves each path of the input column into the destination folder and reports the result per file."
    in_ports = [Port(TABLE, "Source paths")]
    out_ports = [Port(TABLE, "Transfer status")]
    fields = [
        F("column", "column", "Source path column", None, kinds=["string"]),
        F("destination", "string", "Destination folder", "", PATH_HELP, path=True),
        F("mode", "enum", "Transfer", "Copy", options=["Copy", "Move"]),
        F("if_exists", "enum", "If destination exists", "Overwrite", options=["Overwrite", "Ignore", "Fail"]),
    ]

    def configure(self, in_specs, s, ctx):
        require_column(s, "column", in_specs[0], "source path column")
        ctx.resolve_path(s["destination"])
        return [TRANSFER_SCHEMA]

    def execute(self, inputs, s, ctx):
        df = inputs[0]
        col = require_column(s, "column", df.schema, "source path column")
        dest_dir = ctx.resolve_path(s["destination"])
        os.makedirs(dest_dir, exist_ok=True)
        rows = []
        for v in df[col].to_list():
            ctx.check_canceled()
            if not v:
                rows.append((v, None, "missing path"))
                continue
            src = ctx.resolve_path(v)
            dest = os.path.join(dest_dir, os.path.basename(src.rstrip(os.sep)))
            if not os.path.exists(src):
                rows.append((src, dest, "source not found"))
                continue
            if os.path.exists(dest):
                if s["if_exists"] == "Fail":
                    raise InvalidSettingsError(f"Destination exists: {dest}")
                if s["if_exists"] == "Ignore":
                    rows.append((src, dest, "unmodified"))
                    continue
                if os.path.isdir(dest):
                    shutil.rmtree(dest)
                else:
                    os.remove(dest)
            if s["mode"] == "Move":
                shutil.move(src, dest)
                status = "moved"
            else:
                (shutil.copytree if os.path.isdir(src) else shutil.copy2)(src, dest)
                status = "copied"
            rows.append((src, dest, status))
        cols = list(zip(*rows)) if rows else [[], [], []]
        return [pl.DataFrame({n: pl.Series(n, list(v), dtype=pl.String) for n, v in zip(TRANSFER_SCHEMA, cols)})]


class CompressFiles(NodeModel):
    factory = "verknode.io.CompressFiles"
    name = "Compress Files/Folder"
    category = FH
    node_type = "Sink"
    icon = "ZIP"
    short_description = "Creates a zip or tar.gz archive."
    description = (
        "Compresses a single file or folder, or all paths of a column of the (optional) input table, into an "
        "archive. Folders are added with their content. The archive path is pushed as flow variable "
        "<code>archive_path</code>."
    )
    in_ports = [Port(TABLE, "Paths", "Table with a path column", optional=True)]
    out_ports = []
    fields = [
        F("source", "string", "File or folder to compress", "", "Used when no input table is connected.", path=True),
        F("column", "column", "Path column", None, "Used when the input table is connected.", kinds=["string"]),
        F("archive", "string", "Archive file", "archive.zip", PATH_HELP, path=True),
        F("format", "enum", "Format", "zip", options=["zip", "tar.gz"]),
        F("if_exists", "enum", "If archive exists", "Overwrite", options=IF_EXISTS),
    ]

    def configure(self, in_specs, s, ctx):
        if in_specs[0] is not None:
            require_column(s, "column", in_specs[0], "path column")
        elif not s["source"]:
            raise InvalidSettingsError("No file or folder to compress")
        ctx.push_variable("archive_path", ctx.resolve_path(s["archive"]), "STRING")
        return []

    def execute(self, inputs, s, ctx):
        df = inputs[0]
        if df is not None:
            col = require_column(s, "column", df.schema, "path column")
            sources = [ctx.resolve_path(v) for v in df[col].to_list() if v]
        else:
            sources = [ctx.resolve_path(s["source"])]
        for src in sources:
            if not os.path.exists(src):
                raise InvalidSettingsError(f"File does not exist: {src}")
        archive = _check_target(ctx, s["archive"], s["if_exists"])
        _mkparents(archive)
        n = 0
        if s["format"] == "zip":
            with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as z:
                for src in sources:
                    base = os.path.dirname(src.rstrip(os.sep))
                    walk = os.walk(src) if os.path.isdir(src) else [(base, [], [os.path.basename(src)])]
                    for dirpath, _, files in walk:
                        for f in files:
                            ctx.check_canceled()
                            full = os.path.join(dirpath, f)
                            if os.path.abspath(full) == os.path.abspath(archive):
                                continue
                            z.write(full, os.path.relpath(full, base))
                            n += 1
        else:
            with tarfile.open(archive, "w:gz") as t:
                for src in sources:
                    t.add(src, arcname=os.path.basename(src.rstrip(os.sep)))
                    n += 1
        ctx.push_variable("archive_path", archive, "STRING")
        ctx.log(f"Compressed {n} entries into {archive}")
        return []


DECOMPRESS_SCHEMA = pl.Schema({"Path": pl.String, "Is directory": pl.Boolean})


def _inside(root: str, path: str) -> bool:
    root = os.path.realpath(root)
    return os.path.realpath(path) == root or os.path.realpath(path).startswith(root + os.sep)


class DecompressFiles(NodeModel):
    factory = "verknode.io.DecompressFiles"
    name = "Decompress Files"
    category = FH
    node_type = "Source"
    icon = "UNZ"
    short_description = "Extracts a zip, tar, tar.gz or gz archive."
    description = "Extracts an archive into the destination folder and lists the extracted files and folders."
    out_ports = [Port(TABLE, "Extracted files")]
    fields = [
        F("archive", "string", "Archive file", "", PATH_HELP, path=True),
        F("destination", "string", "Destination folder", "", PATH_HELP, path=True),
        F("if_exists", "enum", "If file exists", "Overwrite", options=["Overwrite", "Ignore", "Fail"]),
    ]

    def configure(self, in_specs, s, ctx):
        _existing(ctx, s["archive"])
        ctx.resolve_path(s["destination"])
        return [DECOMPRESS_SCHEMA]

    def execute(self, inputs, s, ctx):
        archive = _existing(ctx, s["archive"])
        dest = ctx.resolve_path(s["destination"])
        os.makedirs(dest, exist_ok=True)
        rows = []

        def target(name):
            out = os.path.join(dest, name)
            if not _inside(dest, out):
                raise InvalidSettingsError(f"Archive entry outside destination folder: {name}")
            return out

        def writable(out):
            if os.path.exists(out) and not os.path.isdir(out):
                if s["if_exists"] == "Fail":
                    raise InvalidSettingsError(f"File exists: {out}")
                return s["if_exists"] == "Overwrite"
            return True

        if zipfile.is_zipfile(archive):
            with zipfile.ZipFile(archive) as z:
                for info in z.infolist():
                    ctx.check_canceled()
                    out = target(info.filename)
                    if info.is_dir():
                        os.makedirs(out, exist_ok=True)
                    elif writable(out):
                        os.makedirs(os.path.dirname(out), exist_ok=True)
                        with z.open(info) as src, open(out, "wb") as dst:
                            shutil.copyfileobj(src, dst)
                    rows.append((out.rstrip(os.sep), info.is_dir()))
        elif tarfile.is_tarfile(archive):
            with tarfile.open(archive) as t:
                for m in t.getmembers():
                    ctx.check_canceled()
                    out = target(m.name)
                    if m.isdir():
                        os.makedirs(out, exist_ok=True)
                    elif m.isfile() and writable(out):
                        t.extract(m, dest, filter="data")
                    else:
                        continue
                    rows.append((out, m.isdir()))
        elif archive.endswith(".gz"):
            out = target(os.path.basename(archive)[:-3])
            if writable(out):
                with gzip.open(archive, "rb") as src, open(out, "wb") as dst:
                    shutil.copyfileobj(src, dst)
            rows.append((out, False))
        else:
            raise InvalidSettingsError("Unsupported archive format (zip, tar, tar.gz, gz)")
        cols = list(zip(*rows)) if rows else [[], []]
        return [pl.DataFrame({"Path": pl.Series(list(cols[0]), dtype=pl.String),
                              "Is directory": pl.Series(list(cols[1]), dtype=pl.Boolean)})]


class CreateFileFolderVariables(NodeModel):
    factory = "verknode.io.CreateFileFolderVariables"
    name = "Create File/Folder Variables"
    category = FH
    node_type = "Other"
    icon = "V📁"
    short_description = "Creates path flow variables relative to a base folder."
    description = (
        "Creates one String flow variable per entry, holding the absolute path of the given file or folder "
        "relative to the base folder (default: the folder containing the workflow)."
    )
    out_ports = [Port(FLOWVARIABLE, "Variables connection")]
    fields = [
        F("base", "string", "Base folder", ".", PATH_HELP, path=True),
        F("variables", "records", "Variables", [{"name": "file_path", "path": "data.csv"}], fields=[
            F("name", "string", "Variable name", ""),
            F("path", "string", "File/folder (relative to base)", ""),
        ]),
    ]

    def execute(self, inputs, s, ctx):
        base = ctx.resolve_path(s["base"] or ".")
        seen = set()
        for r in s["variables"]:
            name = (r.get("name") or "").strip()
            if not name:
                continue
            if name in seen:
                raise InvalidSettingsError(f"Duplicate variable name: {name}")
            seen.add(name)
            ctx.push_variable(name, os.path.normpath(os.path.join(base, r.get("path") or "")), "STRING")
        return [FLOWVAR]
